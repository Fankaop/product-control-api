from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from data.models.batch import Batch
from data.models.product import Product

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2D6A9F", end_color="2D6A9F", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center")


def _style_header_row(ws, columns: list[str]) -> None:
    ws.append(columns)
    for cell in ws[ws.max_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def generate_batch_excel(batch: Batch, products: list[Product]) -> str:
    wb = Workbook()

    # --- Лист 1: Информация о партии ---
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Информация о партии"

    info_rows = [
        ("Номер партии", batch.batch_number),
        ("Дата партии", str(batch.batch_date)),
        ("Статус", "Закрыта" if batch.is_closed else "Открыта"),
        ("Описание задания", batch.task_description),
        ("Номенклатура", batch.nomenclature),
        ("Код ЕКН", batch.ekn_code),
        ("Смена", batch.shift),
        ("Бригада", batch.team),
        ("Начало смены", str(batch.shift_start)),
        ("Окончание смены", str(batch.shift_end)),
    ]
    for label, value in info_rows:
        ws1.append([label, value])
        ws1.cell(row=ws1.max_row, column=1).font = Font(bold=True)

    ws1.column_dimensions["A"].width = 25
    ws1.column_dimensions["B"].width = 35

    # --- Лист 2: Продукция ---
    ws2 = wb.create_sheet("Продукция")
    _style_header_row(ws2, ["ID", "Уникальный код", "Агрегирована", "Дата агрегации"])

    for product in products:
        ws2.append([
            product.id,
            product.unique_code,
            "Да" if product.is_aggregated else "Нет",
            str(product.aggregated_at) if product.aggregated_at else "-",
        ])

    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 25
    ws2.column_dimensions["C"].width = 15
    ws2.column_dimensions["D"].width = 25

    # --- Лист 3: Статистика ---
    ws3 = wb.create_sheet("Статистика")

    total = len(products)
    aggregated = sum(1 for p in products if p.is_aggregated)
    remaining = total - aggregated
    rate = round(aggregated / total * 100, 1) if total > 0 else 0.0

    stat_rows = [
        ("Всего продукции", total),
        ("Агрегировано", aggregated),
        ("Осталось", remaining),
        ("Процент выполнения", f"{rate}%"),
    ]
    for label, value in stat_rows:
        ws3.append([label, value])
        ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True)

    ws3.column_dimensions["A"].width = 25
    ws3.column_dimensions["B"].width = 15

    path = f"/tmp/batch_{batch.id}_report.xlsx"
    wb.save(path)
    return path

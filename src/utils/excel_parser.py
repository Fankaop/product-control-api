import csv
from datetime import date, datetime

from openpyxl import load_workbook

# Маппинг колонок из файла → поля BatchCreate
COLUMN_MAP = {
    "НомерПартии": "batch_number",
    "ДатаПартии": "batch_date",
    "Номенклатура": "nomenclature",
    "КодЕКН": "ekn_code",
    "РабочийЦентр": "work_center_name",
    "ИдентификаторРЦ": "work_center_identifier",
    "Смена": "shift",
    "Бригада": "team",
    "ОписаниеЗадания": "task_description",
    "ДатаВремяНачалаСмены": "shift_start",
    "ДатаВремяОкончанияСмены": "shift_end",
}


def _parse_row(headers: list[str], row: list) -> dict:
    raw = {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}
    result = {}
    for col, field in COLUMN_MAP.items():
        value = raw.get(col)
        if value is None:
            continue
        if field == "batch_number":
            result[field] = int(value)
        elif field == "batch_date":
            result[field] = value if isinstance(value, date) else date.fromisoformat(str(value))
        elif field in ("shift_start", "shift_end"):
            result[field] = value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
        else:
            result[field] = str(value).strip()
    return result


def parse_excel(file_path: str) -> tuple[list[dict], list[dict]]:
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active
    assert ws is not None

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []

    headers = [str(h).strip() for h in rows[0]]
    records, errors = [], []

    for i, row in enumerate(rows[1:], start=2):
        if all(v is None for v in row):
            continue
        try:
            records.append(_parse_row(headers, list(row)))
        except Exception as e:
            errors.append({"row": i, "error": str(e)})

    return records, errors


def parse_csv(file_path: str) -> tuple[list[dict], list[dict]]:
    records, errors = [], []
    with open(file_path, encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for i, row in enumerate(reader, start=2):
            try:
                records.append(_parse_row(list(row.keys()), list(row.values())))
            except Exception as e:
                errors.append({"row": i, "error": str(e)})
    return records, errors

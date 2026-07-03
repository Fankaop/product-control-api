import asyncio
import csv
import os
from datetime import datetime, timezone

from celery_app import celery_app
from core.database import async_session_maker
from data.repositories.batch_repository import BatchRepository
from storage.minio_service import minio_service


@celery_app.task(bind=True, max_retries=3)
def export_batches_to_file(self, filters: dict, format: str = "excel"):
    return asyncio.run(_export(self, filters, format))


async def _export(task, filters: dict, format: str):
    async with async_session_maker() as session:
        repo = BatchRepository(session)
        batches = await repo.get_many(**filters)

    if not batches:
        return {"success": False, "error": "No batches found for given filters"}

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")

    try:
        if format == "csv":
            file_path, object_name = _write_csv(batches, timestamp)
        else:
            file_path, object_name = _write_excel(batches, timestamp)

        minio_service.ensure_buckets()
        file_url = minio_service.upload_file(
            bucket="exports",
            file_path=file_path,
            object_name=object_name,
            expires_days=7,
        )
        file_size = os.path.getsize(file_path)
        os.remove(file_path)

        return {
            "success": True,
            "file_url": file_url,
            "file_name": object_name,
            "file_size": file_size,
            "total_batches": len(batches),
        }

    except Exception as exc:
        raise task.retry(exc=exc)


def _write_csv(batches, timestamp: str) -> tuple[str, str]:
    object_name = f"batches_export_{timestamp}.csv"
    file_path = f"/tmp/{object_name}"

    with open(file_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([
            "ID", "НомерПартии", "ДатаПартии", "Статус",
            "Номенклатура", "КодЕКН", "Смена", "Бригада",
            "НачалоСмены", "ОкончаниеСмены", "СозданоВ",
        ])
        for b in batches:
            writer.writerow([
                b.id, b.batch_number, b.batch_date,
                "Закрыта" if b.is_closed else "Открыта",
                b.nomenclature, b.ekn_code, b.shift, b.team,
                b.shift_start, b.shift_end, b.created_at,
            ])

    return file_path, object_name


def _write_excel(batches, timestamp: str) -> tuple[str, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    object_name = f"batches_export_{timestamp}.xlsx"
    file_path = f"/tmp/{object_name}"

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Партии"

    headers = [
        "ID", "НомерПартии", "ДатаПартии", "Статус",
        "Номенклатура", "КодЕКН", "Смена", "Бригада",
        "НачалоСмены", "ОкончаниеСмены", "СозданоВ",
    ]

    ws.append(headers)
    header_fill = PatternFill(start_color="2D6A9F", end_color="2D6A9F", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for b in batches:
        ws.append([
            b.id, b.batch_number, str(b.batch_date),
            "Закрыта" if b.is_closed else "Открыта",
            b.nomenclature, b.ekn_code, b.shift, b.team,
            str(b.shift_start), str(b.shift_end), str(b.created_at),
        ])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20

    wb.save(file_path)
    return file_path, object_name
